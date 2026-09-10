import os
import pandas as pd
from docx import Document
from reportlab.pdfgen import canvas
import smtplib
#Fluxo de automação para gerar arquivos e enviar por e-mail
#exportar Excel
#gerar o Word
#converter para PDF
#Enviar por e-mail

print('1. Gerando o arquivo Excel...')
#Criando os dados para o arquivo Excel
dados_venda = {
    'Produto': ['Produto A', 'Produto B', 'Produto C'],
    'Quantidade': [10, 5, 8],
    'Preço': [100.0, 50.0, 80.0]
}
df = pd.DataFrame(dados_venda)
excel_file = 'vendas.xlsx'
df.to_excel(excel_file, index=False)
print(f'Arquivo Excel "{excel_file}" gerado com sucesso!')

#Ler o excel e gerar o Word
print('2. Criando o relatório no Word...')
df_lido = pd.read_excel(excel_file)
#criando o documento Word
doc = Document()
#Título
doc.add_heading('Relatório Executivo de Vendas', level=1)
#parágrafo
doc.add_paragraph(f'Este relatório foi gerado automaticamente a partir do arquivo \'{excel_file}\'.')

doc.add_heading('Detalhamento por Produto', level=2)
#Tabela
table =